import os
from openpyxl import Workbook
from ortak import *
from ccola_veri import BILANCO, GELIR, NAKIT, YILLAR, PIYASA
OUT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "07-vaka-sirketi-CCOLA"))
KAYNAK = ("Kaynak KAP'ta yayımlanan konsolide finansal tablolar, Fintables veri tabanı üzerinden 05.10.2026 tarihinde alındı. "
          "Tutarlar TMS 29 kapsamında son açıklanan dönemin (2026/06) satın alma gücüyle ifade edilmiştir. "
          "Faaliyet raporundaki tutarlar 2025/12 alım gücündedir ve yaklaşık 1,1776 kat düşüktür, oranlar aynıdır.")
LINKLER = [
    ("2025 yıllık finansal rapor (KAP)", "https://www.kap.org.tr/tr/Bildirim/1565234"),
    ("2025 entegre faaliyet raporu (KAP)", "https://www.kap.org.tr/tr/Bildirim/1565282"),
    ("2025 entegre faaliyet raporu PDF", "https://storage.fintables.com/media/uploads/kap-attachments/CCI%202025%20Entegre%20Faaliyet%20Raporu%20TR.pdf"),
    ("2024 yıllık finansal rapor (KAP)", "https://www.kap.org.tr/tr/Bildirim/1400442"),
]

def tablolar():
    wb = Workbook()
    wk = wb.active
    wk.title = "Bilgi"
    wk.sheet_view.showGridLines = False
    wk.column_dimensions["A"].width = 40
    wk.column_dimensions["B"].width = 100
    wk["A1"] = "Coca-Cola İçecek A.Ş. (CCOLA) Vaka Dosyası"
    wk["A1"].font = f_baslik
    wk["A2"] = "Bilanço Okuma Eğitimi, Ders 23 vaka şirketi. Yıllık konsolide tablolar, 2023 ile 2025 arası, milyon TL."
    wk["A2"].font = f_alt
    bilgiler = [("Sektör", PIYASA["sektor"]), ("Ödenmiş sermaye (pay adedi)", f"{PIYASA['pay_adedi']:,}".replace(",", ".")),
                ("Fiili dolaşım oranı", f"%{PIYASA['fiili_dolasim']:.2f}".replace(".", ",")),
                ("Raporlama esası", "TFRS, konsolide, TMS 29 uygulanmış"), ("Para birimi", "milyon TL")]
    for i, (a, b) in enumerate(bilgiler):
        wk.cell(row=4 + i, column=1, value=a).font = f_kalin
        wk.cell(row=4 + i, column=2, value=b).font = f_norm
    wk.cell(row=10, column=1, value="Önemli not").font = f_bolum
    c = wk.cell(row=11, column=1, value=KAYNAK)
    c.font, c.alignment = f_norm, sar
    wk.merge_cells("A11:B13")
    wk.row_dimensions[11].height = 30
    wk.cell(row=15, column=1, value="Orijinal belgeler").font = f_bolum
    for i, (a, u) in enumerate(LINKLER):
        wk.cell(row=16 + i, column=1, value=a).font = f_norm
        h = wk.cell(row=16 + i, column=2, value=u)
        h.hyperlink, h.font = u, Font(name=FONT, size=10, color="0563C1", underline="single")
    wk.cell(row=21, column=1, value="Bu dosyadaki sayfalar").font = f_bolum
    for i, t in enumerate(["Bilanço, Gelir Tablosu ve Nakit Akış sayfaları ham tablolardır.",
                           "Rasyo hesabı ve Altman Beneish hesabı bu klasördeki diğer iki Excel dosyasındadır."]):
        wk.cell(row=22 + i, column=1, value=t).font = f_norm
    for ad, veri in (("Bilanço", BILANCO), ("Gelir Tablosu", GELIR), ("Nakit Akış", NAKIT)):
        ws = wb.create_sheet(ad)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 58
        for col in "BCD":
            ws.column_dimensions[col].width = 15
        ws["A1"] = f"CCOLA {ad} (milyon TL)"
        ws["A1"].font = f_baslik
        ws["A2"] = "Yıllık konsolide, 2026/06 satın alma gücüyle"
        ws["A2"].font = f_alt
        kafa(ws, 4, ["Kalem"] + [f"{y}/12" for y in YILLAR])
        r = 5
        for kalem, deger in veri:
            if deger is None:
                bolum(ws, r, kalem, 4)
            else:
                toplam = kalem.isupper() or kalem.startswith(("Toplam", "A.", "B.", "C.", "Brüt", "Faaliyet Karı", "Dönem Karı", "FAVÖK", "Faaliyetlerden Elde"))
                ws.cell(row=r, column=1, value=kalem).font = f_kalin if toplam else f_norm
                for i, v in enumerate(deger):
                    c = ws.cell(row=r, column=2 + i, value=round(v / 1e6, 1))
                    c.number_format, c.font = NUM, f_kalin if toplam else f_norm
                for col in range(1, 5):
                    ws.cell(row=r, column=col).border = kenar
            r += 1
        ws.cell(row=r + 1, column=1, value=KAYNAK).font = f_alt
        ws.freeze_panes = "B5"
    wb.save(f"{OUT}/CCOLA-finansal-tablolar-2023-2025.xlsx")

if __name__ == "__main__":
    tablolar()
