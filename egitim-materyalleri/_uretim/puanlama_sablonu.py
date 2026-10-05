from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from ortak import *

KATEGORILER = [
    ("Karlılık", "Şirketin kar üretme gücü", [
        "Brüt kar marjı seviyesi ve istikrarı",
        "FAVÖK marjı ve sektör ortalamasıyla kıyas",
        "Net kar marjı ve faaliyet karı ile makas",
        "ROE ve ROIC, enflasyon ve sermaye maliyetine göre",
    ]),
    ("Büyüme", "Gelir ve kar büyüme dinamiği", [
        "Reel hasılat büyümesi (hacim ve fiyat ayrımı)",
        "FAVÖK büyümesi",
        "Net kar büyümesi",
        "Organik ve inorganik büyüme ayrımı",
    ]),
    ("Borçluluk", "Finansal risk ve kaldıraç", [
        "Net borç / FAVÖK",
        "Finansal borç / Özkaynak",
        "Faiz karşılama oranı",
        "Kur riski ve borç vade yapısı",
    ]),
    ("Nakit Üretimi", "Gerçek nakit yaratma kapasitesi", [
        "Serbest nakit akışı trendi",
        "Kar kalitesi (İşletme nakit akışı / Net kar)",
        "Nakit dönüşüm süresi trendi",
        "CAPEX / Amortisman ve yatırım disiplini",
    ]),
    ("Temettü Kapasitesi", "Sürdürülebilir getiri potansiyeli", [
        "Temettü verimi",
        "Dağıtım oranı",
        "Temettü / SNA",
        "Temettü geçmişi ve büyüme trendi",
    ]),
    ("Rekabet Avantajı", "Ekonomik hendek (moat) gücü", [
        "Marka gücü ve fiyatlama gücü",
        "Maliyet avantajı ve ölçek",
        "Geçiş maliyeti ve ağ etkisi",
        "Düzenleyici bariyer ve lisanslar",
    ]),
    ("Yönetim Kalitesi", "Sermaye tahsisi ve şeffaflık", [
        "Sermaye getirisi trendi ve yatırım kararları",
        "İlişkili taraf işlemleri",
        "Azınlık haklarına saygı ve kurumsal yönetim",
        "Açıklanan hedeflere ulaşma ve şeffaflık",
    ]),
    ("Değerleme", "Mevcut fiyatın makullüğü", [
        "F/K, tarihsel ve sektör ortalamasıyla",
        "FD/FAVÖK, tarihsel ve sektör ortalamasıyla",
        "PD/DD ve ROE dengesi",
        "SNA verimi",
    ]),
]

ORNEK = {  # Karlılık kategorisi için doldurulmuş örnek (CCOLA 2025/12)
    0: (8, "Brüt marj 3 yılda %32,7 seviyesinden %35,6 seviyesine çıktı, TMS 29 altında güçlü fiyatlama gücü."),
    1: (7, "FAVÖK marjı %17,5 ile %17,7 arasında çok istikrarlı."),
    2: (5, "Net marj %20,4 seviyesinden %7,5 seviyesine indi, fark parasal kazancın azalmasından geliyor."),
    3: (6, "ROIC %14,8, ROE %19,3. Pozitif ama düşüş trendinde."),
}


def build(path, sirket="", donem="", ornek=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Puanlama"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", [22, 46, 11, 70, 14]):
        ws.column_dimensions[col].width = w
    ws["A1"] = "8 Kategorili Puanlama Şablonu"
    ws["A1"].font = f_baslik
    ws["A2"] = "Bilanço Okuma Eğitimi, ek materyal 4. Her ölçüte 1 ile 10 arası tam puan verin, gerekçeyi veriye dayandırın."
    ws["A2"].font = f_alt
    ws["A3"], ws["A4"] = "Şirket", "Dönem"
    ws["A3"].font = ws["A4"].font = f_kalin
    for ref, v in (("B3", sirket), ("B4", donem)):
        ws[ref] = v or None
        ws[ref].font, ws[ref].fill, ws[ref].border = f_girdi, fill_girdi, kenar
    kafa(ws, 6, ["Kategori", "Ölçüt", "Puan (1-10)", "Gerekçe (somut veri ile)", "Kategori puanı"])
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True,
                        showErrorMessage=True, errorTitle="Geçersiz puan", error="1 ile 10 arasında tam sayı girin")
    ws.add_data_validation(dv)
    r = 7
    kat_hucre = []
    for ki, (kat, olcer, olcutler) in enumerate(KATEGORILER):
        bas = r
        for oi, o in enumerate(olcutler):
            ws.cell(row=r, column=2, value=o).font = f_norm
            p = ws.cell(row=r, column=3)
            g = ws.cell(row=r, column=4)
            if ornek and ki == 0:
                p.value, g.value = ORNEK[oi]
            p.font, p.fill, p.border, p.alignment = f_girdi, fill_girdi, kenar, orta
            g.font, g.fill, g.border, g.alignment = f_girdi, fill_girdi, kenar, sar
            dv.add(p)
            ws.cell(row=r, column=2).border = kenar
            ws.cell(row=r, column=2).alignment = sar
            ws.row_dimensions[r].height = 30
            r += 1
        son = r - 1
        ws.merge_cells(start_row=bas, start_column=1, end_row=son, end_column=1)
        a = ws.cell(row=bas, column=1, value=f"{kat}\n{olcer}")
        a.font, a.alignment = f_kalin, sar
        for rr in range(bas, son + 1):
            ws.cell(row=rr, column=1).border = kenar
            ws.cell(row=rr, column=5).border = kenar
        ws.merge_cells(start_row=bas, start_column=5, end_row=son, end_column=5)
        e = ws.cell(row=bas, column=5, value=f'=IF(COUNT(C{bas}:C{son})=0,"",ROUND(AVERAGE(C{bas}:C{son}),1))')
        e.font, e.alignment, e.fill, e.number_format = f_kalin, orta, fill_sonuc, "0.0"
        kat_hucre.append((kat, f"E{bas}"))
    # ---- Özet sayfası
    wo = wb.create_sheet("Özet", 0)
    wo.sheet_view.showGridLines = False
    for col, w in zip("ABCD", [28, 16, 70, 4]):
        wo.column_dimensions[col].width = w
    wo["A1"] = "Puanlama Özeti"
    wo["A1"].font = f_baslik
    wo["A2"] = "=IF(Puanlama!B3=\"\",\"Şirket adını Puanlama sayfasına yazın\",Puanlama!B3&\"  \"&Puanlama!B4)"
    wo["A2"].font = Font(name=FONT, size=11, bold=True, color="008000")
    kafa(wo, 4, ["Kategori", "Puan (1-10)", "Ne ölçer"])
    for i, (kat, ref) in enumerate(kat_hucre):
        rr = 5 + i
        wo.cell(row=rr, column=1, value=kat).font = f_norm
        c = wo.cell(row=rr, column=2, value=f"=Puanlama!{ref}")
        c.font, c.number_format, c.alignment = f_link, "0.0", orta
        wo.cell(row=rr, column=3, value=KATEGORILER[i][1]).font = f_norm
        for col in range(1, 4):
            wo.cell(row=rr, column=col).border = kenar
    son = 5 + len(kat_hucre) - 1
    satirlar = [
        ("Puanlanan kategori sayısı", f"=COUNT(B5:B{son})", "0", "Veri yetersiz kategoriyi boş bırakın, genel puan kalan kategoriler üzerinden hesaplanır."),
        ("Toplam puan", f'=IF(B{son+2}=0,"",SUM(B5:B{son}))', "0.0", None),
        ("Azami puan", f"=B{son+2}*10", "0", "8 kategori puanlandıysa 80."),
        ("Genel puan", f'=IF(B{son+2}=0,"",B{son+3}/B{son+4})', "0.0%", None),
        ("Yorum bandı", f'=IF(B{son+5}="","",IF(B{son+5}>=0.8,"Güçlü şirket, temel göstergeler büyük ölçüde olumlu",IF(B{son+5}>=0.6,"Sağlam temel, bazı alanlarda iyileşme potansiyeli var",IF(B{son+5}>=0.4,"Karışık tablo, belirgin zayıflıklar mevcut",IF(B{son+5}>=0.2,"Zayıf temel, yüksek riskli","Ciddi sorunlu, temel göstergeler büyük ölçüde olumsuz")))))', None, None),
    ]
    for i, (ad, f, fmt, notu) in enumerate(satirlar):
        rr = son + 2 + i
        wo.cell(row=rr, column=1, value=ad).font = f_kalin
        c = wo.cell(row=rr, column=2, value=f)
        c.font, c.fill, c.border = f_kalin, fill_sonuc, kenar
        if fmt:
            c.number_format = fmt
            c.alignment = orta
        if notu:
            wo.cell(row=rr, column=3, value=notu).font = f_alt
    wo.merge_cells(start_row=son + 6, start_column=2, end_row=son + 6, end_column=3)
    wo.cell(row=son + 6, column=2).alignment = Alignment(vertical="center")
    wo.conditional_formatting.add(f"B5:B{son}", CellIsRule(operator="greaterThanOrEqual", formula=["7"], fill=PatternFill("solid", fgColor="C6EFCE")))
    wo.conditional_formatting.add(f"B5:B{son}", CellIsRule(operator="between", formula=["0.1", "4.9"], fill=PatternFill("solid", fgColor="FFC7CE")))
    rr = son + 9
    wo.cell(row=rr, column=1, value="Puan ölçeği").font = f_bolum
    olcek = [("9-10", "Olağanüstü, sektörün en iyileri arasında"), ("7-8", "Güçlü, ortalamanın belirgin üzerinde"),
             ("5-6", "Ortalama, kabul edilebilir ama öne çıkmıyor"), ("3-4", "Ortalamanın altında, dikkat gerektiriyor"),
             ("1-2", "Zayıf, ciddi endişe kaynağı")]
    for i, (p, a) in enumerate(olcek):
        wo.cell(row=rr + 1 + i, column=1, value=p).font = f_kalin
        wo.cell(row=rr + 1 + i, column=2, value=a).font = f_norm
    rr += 7
    wo.cell(row=rr, column=1, value="Puanlama kuralları").font = f_bolum
    kurallar = [
        "Her puan somut veriye dayanmalı. Gerekçe sütununa rakamı ve kaynağını yazın.",
        "Puanları sektör bağlamında verin. Enerji şirketinin borçluluk normu ile yazılım şirketininki farklıdır.",
        "Değerleme puanı diğerlerinden bağımsızdır. Çok iyi bir şirket pahalıysa değerleme puanı düşük olabilir.",
        "Kategori puanı, doldurulan ölçütlerin ortalamasıdır. Boş bıraktığınız ölçüt hesaba girmez.",
        "Genel puan bir finansal sağlık göstergesidir, yatırım tavsiyesi değildir.",
    ]
    for i, k in enumerate(kurallar):
        c = wo.cell(row=rr + 1 + i, column=1, value=f"{i+1}. {k}")
        c.font = f_norm
        wo.merge_cells(start_row=rr + 1 + i, start_column=1, end_row=rr + 1 + i, end_column=3)
    if ornek:
        ws.cell(row=r + 1, column=1, value="Karlılık satırları CCOLA 2025/12 verisiyle örnek olarak doldurulmuştur. Kendi analizinizde üzerine yazın.").font = f_alt
    lejant(wo, rr + 8)
    wb.save(path)
