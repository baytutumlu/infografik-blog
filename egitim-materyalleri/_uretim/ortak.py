from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

FONT = "Arial"
LACIVERT = "1F3A5F"
f_baslik = Font(name=FONT, size=14, bold=True, color=LACIVERT)
f_alt = Font(name=FONT, size=10, italic=True, color="555555")
f_kafa = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_bolum = Font(name=FONT, size=10, bold=True, color=LACIVERT)
f_norm = Font(name=FONT, size=10)
f_kalin = Font(name=FONT, size=10, bold=True)
f_girdi = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
fill_kafa = PatternFill("solid", fgColor=LACIVERT)
fill_bolum = PatternFill("solid", fgColor="DCE6F0")
fill_girdi = PatternFill("solid", fgColor="FFF2CC")
fill_sonuc = PatternFill("solid", fgColor="E2EFDA")
ince = Side(style="thin", color="BFBFBF")
kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
sar = Alignment(wrap_text=True, vertical="top")
orta = Alignment(horizontal="center", vertical="center", wrap_text=True)

NUM = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
KAT = '0.00"x";(0.00"x");"-"'
GUN = '#,##0" gün";(#,##0" gün");"-"'
DEC = '0.000;(0.000);"-"'

def kafa(ws, row, basliklar, col=1):
    for i, b in enumerate(basliklar):
        c = ws.cell(row=row, column=col + i, value=b)
        c.font, c.fill, c.alignment, c.border = f_kafa, fill_kafa, orta, kenar

def bolum(ws, row, metin, ncol):
    for i in range(1, ncol + 1):
        c = ws.cell(row=row, column=i)
        c.fill = fill_bolum
        c.border = kenar
    ws.cell(row=row, column=1, value=metin).font = f_bolum

def lejant(ws, row, col=1):
    ws.cell(row=row, column=col, value="Renk kodları").font = f_kalin
    a = ws.cell(row=row + 1, column=col, value="Mavi yazı, sarı zemin")
    a.font, a.fill = f_girdi, fill_girdi
    ws.cell(row=row + 1, column=col + 1, value="Sizin dolduracağınız girdi hücresi").font = f_norm
    b = ws.cell(row=row + 2, column=col, value="Siyah yazı")
    b.font = f_norm
    ws.cell(row=row + 2, column=col + 1, value="Formül, elle değiştirmeyin").font = f_norm
    g = ws.cell(row=row + 3, column=col, value="Yeşil yazı")
    g.font = f_link
    ws.cell(row=row + 3, column=col + 1, value="Başka sayfadan gelen bağlantı").font = f_norm
