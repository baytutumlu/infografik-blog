from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from ortak import *

# (anahtar, görünen ad, finansal tablodaki karşılığı)
GIRDILER = [
    ("BİLANÇO", None, None),
    ("nakit", "Nakit ve nakit benzerleri", "Bilanço, Dönen Varlıklar, Nakit ve Nakit Benzerleri"),
    ("fin_yat", "Kısa vadeli finansal yatırımlar", "Bilanço, Dönen Varlıklar, Finansal Yatırımlar"),
    ("tic_alacak", "Ticari alacaklar (KV)", "Bilanço, Dönen Varlıklar, Ticari Alacaklar"),
    ("stok", "Stoklar", "Bilanço, Dönen Varlıklar, Stoklar"),
    ("donen", "Toplam dönen varlıklar", "Bilanço, Toplam Dönen Varlıklar"),
    ("mdv", "Maddi duran varlıklar", "Bilanço, Duran Varlıklar, Maddi Duran Varlıklar"),
    ("tv", "Toplam varlıklar", "Bilanço, Toplam Varlıklar (Aktif toplamı)"),
    ("kv_fin", "Kısa vadeli finansal borçlar", "Bilanço, KV Yükümlülükler, Finansal Borçlar (UV borçların KV kısımları dahil)"),
    ("tic_borc", "Ticari borçlar (KV)", "Bilanço, KV Yükümlülükler, Ticari Borçlar"),
    ("kvyk", "Toplam kısa vadeli yükümlülükler", "Bilanço, Toplam Kısa Vadeli Yükümlülükler"),
    ("uv_fin", "Uzun vadeli finansal borçlar", "Bilanço, UV Yükümlülükler, Finansal Borçlar"),
    ("uvyk", "Toplam uzun vadeli yükümlülükler", "Bilanço, Toplam Uzun Vadeli Yükümlülükler"),
    ("ana_oz", "Ana ortaklığa ait özkaynaklar", "Bilanço, Özkaynaklar, Ana Ortaklığa Ait Özkaynaklar"),
    ("top_oz", "Toplam özkaynaklar", "Bilanço, Toplam Özkaynaklar (azınlık payları dahil)"),
    ("GELİR TABLOSU (giderleri pozitif girin)", None, None),
    ("hasilat", "Hasılat", "Gelir Tablosu, Hasılat (Satış Gelirleri)"),
    ("smm", "Satışların maliyeti", "Gelir Tablosu, Satışların Maliyeti (-)"),
    ("gyg", "Genel yönetim giderleri", "Gelir Tablosu, Genel Yönetim Giderleri (-)"),
    ("pazarlama", "Pazarlama, satış ve dağıtım giderleri", "Gelir Tablosu, Pazarlama Giderleri (-)"),
    ("arge", "Araştırma ve geliştirme giderleri", "Gelir Tablosu, Araştırma ve Geliştirme Giderleri (-), yoksa 0"),
    ("faal_kar", "Esas faaliyet karı", "Gelir Tablosu, Esas Faaliyet Karı (Zararı)"),
    ("fvok", "Finansman öncesi faaliyet karı (FVÖK)", "Gelir Tablosu, Finansman Geliri (Gideri) Öncesi Faaliyet Karı"),
    ("amort", "Amortisman ve itfa payları", "Nakit Akış Tablosu, Amortisman ve İtfa Gideri ile İlgili Düzeltmeler"),
    ("fin_gider", "Finansman giderleri", "Gelir Tablosu, Finansman Giderleri (-)"),
    ("vok", "Vergi öncesi kar", "Gelir Tablosu, Sürdürülen Faaliyetler Vergi Öncesi Karı"),
    ("vergi", "Vergi gideri (dönem + ertelenmiş)", "Gelir Tablosu, Sürdürülen Faaliyetler Vergi Gideri (-)"),
    ("net_kar", "Dönem karı (azınlık dahil)", "Gelir Tablosu, Dönem Karı (Zararı)"),
    ("ana_net", "Ana ortaklık payı net kar", "Gelir Tablosu, Dönem Karının Dağılımı, Ana Ortaklık Payları"),
    ("NAKİT AKIŞ TABLOSU (çıkışları pozitif girin)", None, None),
    ("ocf", "İşletme faaliyetlerinden nakit akışı", "Nakit Akış Tablosu, A. İşletme Faaliyetlerinden Nakit Akışları"),
    ("capex", "Yatırım harcamaları (CAPEX)", "Nakit Akış Tablosu, Maddi ve Maddi Olmayan Duran Varlık Alımından Nakit Çıkışları"),
    ("temettu", "Ödenen temettü", "Nakit Akış Tablosu, Finansman Faaliyetleri, Ödenen Temettüler"),
]

PER_COLS = ["C", "D", "E"]  # Girdi ve Rasyolar sayfasında dönem sütunları


def build(path, donemler=None, veri=None, ornek=None, baslik_ek=""):
    wb = Workbook()
    # ---------------- Nasıl Kullanılır ----------------
    ws0 = wb.active
    ws0.title = "Nasıl Kullanılır"
    ws0.sheet_view.showGridLines = False
    ws0.column_dimensions["A"].width = 4
    ws0.column_dimensions["B"].width = 100
    ws0["B2"] = "Rasyo Hesap Şablonu" + baslik_ek
    ws0["B2"].font = f_baslik
    ws0["B3"] = "Bilanço Okuma Eğitimi, ek materyal 1"
    ws0["B3"].font = f_alt
    adimlar = [
        "1. Girdi sayfasına geçin. Üst satıra dönem adlarını yazın (örnek 2023/12, 2024/12, 2025/12). En eski dönem solda olsun.",
        "2. Sadece sarı zeminli mavi hücreleri doldurun. Her satırın yanında rakamı tablonun hangi kaleminden alacağınız yazıyor.",
        "3. Tüm tutarları aynı birimde girin. Şablon milyon TL varsayıyor, bin TL kullanacaksanız hepsini bin TL girin, oranlar değişmez.",
        "4. Giderleri, vergiyi, finansman giderini, CAPEX ve temettüyü pozitif sayı olarak girin. Formüller işareti kendisi düzeltiyor.",
        "5. Rasyolar sayfası otomatik hesaplanır. Ortalama gerektiren oranlar (devir hızları, ROE, ROA, ROIC) bir önceki dönem doluysa ortalama bakiyeyi, boşsa dönem sonu bakiyeyi kullanır.",
        "6. Tek dönemle de çalışabilirsiniz ama trend görmek için en az iki, tercihen üç yıl girin.",
        "7. Yorum sütunundaki eşikler genel referanstır. Perakende, enerji, holding ve banka gibi sektörlerde normlar çok farklıdır, kararınızı sektör ortalamasıyla karşılaştırarak verin.",
        "8. Yüksek enflasyon muhasebesi (TMS 29) uygulayan şirketlerde tüm dönemlerin aynı satın alma gücüyle raporlandığından emin olun. Önceki yıl rakamlarını her zaman son raporun karşılaştırmalı sütunundan alın.",
        "9. Kontrol sayfası bilanço denkliğini ve temel tutarlılıkları kontrol eder. Orada UYARI görürseniz girdileri yeniden gözden geçirin.",
        "",
        "Bu dosya eğitim amaçlıdır, yatırım tavsiyesi değildir.",
    ]
    for i, a in enumerate(adimlar):
        c = ws0.cell(row=5 + i, column=2, value=a)
        c.font, c.alignment = f_norm, sar
        ws0.row_dimensions[5 + i].height = 30 if len(a) > 95 else 16
    lejant(ws0, 18, col=2)
    ws0.column_dimensions["C"].width = 40

    # ---------------- Girdi ----------------
    ws = wb.create_sheet("Girdi")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [42, 62, 14, 14, 14, 3, 22]):
        ws.column_dimensions[col].width = w
    ws["A1"] = "Girdi Tablosu (milyon TL)" + baslik_ek
    ws["A1"].font = f_baslik
    ws["A2"] = "Sarı hücreleri doldurun. Giderler ve nakit çıkışları pozitif girilir."
    ws["A2"].font = f_alt
    kafa(ws, 4, ["Kalem", "Finansal tablodaki karşılığı", "Dönem 1", "Dönem 2", "Dönem 3"])
    kafa(ws, 4, ["Örnek (CCOLA 2025/12)"], col=7)
    # dönem adları satırı
    ws["A5"], ws["B5"] = "Dönem adı", "Örnek 2023/12, en eski dönem solda"
    ws["A5"].font = f_kalin
    ws["B5"].font = f_alt
    for i, col in enumerate(PER_COLS):
        c = ws[f"{col}5"]
        c.value = donemler[i] if donemler else None
        c.font, c.fill, c.border, c.alignment = f_girdi, fill_girdi, kenar, orta
    ws["G5"] = "2025/12"
    ws["G5"].font = f_alt
    satir = {}
    r = 6
    for key, ad, kars in GIRDILER:
        if ad is None:
            bolum(ws, r, key, 5)
            r += 1
            continue
        satir[key] = r
        ws.cell(row=r, column=1, value=ad).font = f_norm
        k = ws.cell(row=r, column=2, value=kars)
        k.font, k.alignment = Font(name=FONT, size=9, color="555555"), sar
        for i, col in enumerate(PER_COLS):
            c = ws[f"{col}{r}"]
            c.value = (veri[key][i] if veri and veri[key][i] is not None else None)
            c.font, c.fill, c.border, c.number_format = f_girdi, fill_girdi, kenar, NUM
        if ornek:
            e = ws.cell(row=r, column=7, value=ornek[key])
            e.font, e.number_format = f_alt, NUM
        for col in "AB":
            ws[f"{col}{r}"].border = kenar
        r += 1
    ws.freeze_panes = "C6"
    if veri is not None and "_kaynak" in veri:
        ws.cell(row=r + 1, column=1, value=veri["_kaynak"]).font = f_alt

    # ---------------- Rasyolar ----------------
    wr = wb.create_sheet("Rasyolar")
    wr.sheet_view.showGridLines = False
    for col, w in zip("ABCDEF", [40, 52, 14, 14, 14, 60]):
        wr.column_dimensions[col].width = w
    wr["A1"] = "Rasyolar" + baslik_ek
    wr["A1"].font = f_baslik
    wr["A2"] = "Bu sayfadaki tüm hücreler formüldür. Değiştirmek için Girdi sayfasını kullanın."
    wr["A2"].font = f_alt
    kafa(wr, 4, ["Rasyo", "Formül", "Dönem 1", "Dönem 2", "Dönem 3", "Nasıl yorumlanır"])
    for col in PER_COLS:
        c = wr[f"{col}5"]
        c.value = f'=IF(Girdi!{col}5="","",Girdi!{col}5)'
        c.font, c.alignment = Font(name=FONT, size=10, bold=True, color="008000"), orta
    wr["A5"] = "Dönem"
    wr["A5"].font = f_kalin

    def g(key, col):
        return f"Girdi!{col}{satir[key]}"

    def ort(key, col):
        i = PER_COLS.index(col)
        if i == 0:
            return g(key, col)
        return f"AVERAGE({g(key, col)},{g(key, PER_COLS[i-1])})"

    rsatir = {}

    def ortr(rkey, col):
        i = PER_COLS.index(col)
        here = f"{col}{rsatir[rkey]}"
        if i == 0:
            return here
        return f"AVERAGE({here},{PER_COLS[i-1]}{rsatir[rkey]})"

    favok = lambda c: f"({g('hasilat',c)}-{g('smm',c)}-{g('gyg',c)}-{g('pazarlama',c)}-{g('arge',c)}+{g('amort',c)})"
    netborc = lambda c: f"({g('kv_fin',c)}+{g('uv_fin',c)}-{g('nakit',c)}-{g('fin_yat',c)})"

    RAS = [
        ("LİKİDİTE ORANLARI",),
        ("cari", "Cari oran", "Dönen varlıklar / KV yükümlülükler", lambda c: f"{g('donen',c)}/{g('kvyk',c)}", KAT,
         "1,5 üzeri genelde rahat, 1 altı dikkat ister. Hızlı tüketim ve perakendede 1 altı normal olabilir, çünkü tedarikçi kredisi kullanılır."),
        ("asit", "Asit test oranı", "(Dönen varlıklar - Stoklar) / KV yükümlülükler", lambda c: f"({g('donen',c)}-{g('stok',c)})/{g('kvyk',c)}", KAT,
         "Stokları satmadan KV borç ödeme gücü. 1 civarı sağlıklı kabul edilir."),
        ("nakitoran", "Nakit oranı", "(Nakit + KV finansal yatırımlar) / KV yükümlülükler", lambda c: f"({g('nakit',c)}+{g('fin_yat',c)})/{g('kvyk',c)}", KAT,
         "En katı likidite ölçüsü. 0,2 ile 0,5 arası çoğu sanayi şirketi için yeterli."),
        ("nis_brut", "Net çalışma sermayesi (mn TL)", "Dönen varlıklar - KV yükümlülükler", lambda c: f"{g('donen',c)}-{g('kvyk',c)}", NUM,
         "Negatifse şirket KV kaynakla UV varlık finanse ediyor olabilir. Trendine bakın."),
        ("BORÇLULUK ORANLARI",),
        ("kaldirac", "Kaldıraç oranı", "Toplam yükümlülükler / Toplam varlıklar", lambda c: f"({g('kvyk',c)}+{g('uvyk',c)})/{g('tv',c)}", PCT,
         "Varlıkların ne kadarı borçla finanse edilmiş. %60 üzeri sanayide yüksek sayılır."),
        ("oz_aktif", "Özkaynak / Aktif", "Toplam özkaynaklar / Toplam varlıklar", lambda c: f"{g('top_oz',c)}/{g('tv',c)}", PCT,
         "Kaldıraç oranının tamamlayıcısı. Yükseldikçe finansal risk azalır."),
        ("borc_oz", "Finansal borç / Özkaynak", "(KV + UV finansal borçlar) / Toplam özkaynaklar", lambda c: f"({g('kv_fin',c)}+{g('uv_fin',c)})/{g('top_oz',c)}", KAT,
         "1 üzeri ise şirket ortaktan çok bankayla finanse ediliyor demektir."),
        ("netborc", "Net borç (mn TL)", "KV + UV finansal borçlar - Nakit - KV finansal yatırımlar", lambda c: netborc(c), NUM,
         "Negatifse şirket net nakit pozisyonunda."),
        ("nb_favok", "Net borç / FAVÖK", "Net borç / FAVÖK", lambda c: f"{netborc(c)}/{favok(c)}", KAT,
         "2 altı rahat, 2 ile 3,5 arası izlenmeli, 3,5 üzeri yüksek. Negatif değer net nakit demek."),
        ("faiz_kars", "Faiz karşılama oranı", "FVÖK / Finansman giderleri", lambda c: f"{g('fvok',c)}/{g('fin_gider',c)}", KAT,
         "3 üzeri rahat, 1,5 altı alarm. Finansman giderinde kur farkı varsa ayrıca dipnottan faiz giderine bakın."),
        ("kv_pay", "KV finansal borç payı", "KV finansal borç / Toplam finansal borç", lambda c: f"{g('kv_fin',c)}/({g('kv_fin',c)}+{g('uv_fin',c)})", PCT,
         "Yüksekse yeniden finansman riski artar. Nakit oranıyla birlikte okuyun."),
        ("FAALİYET (VERİMLİLİK) ORANLARI",),
        ("aktif_devir", "Aktif devir hızı", "Hasılat / Ortalama toplam varlıklar", lambda c: f"{g('hasilat',c)}/{ort('tv',c)}", KAT,
         "Her 1 TL varlıkla ne kadar satış yapıldığı. Sermaye yoğun sektörlerde düşük olur."),
        ("mdv_devir", "Maddi duran varlık devir hızı", "Hasılat / Ortalama maddi duran varlıklar", lambda c: f"{g('hasilat',c)}/{ort('mdv',c)}", KAT,
         "Düşüşü yeni yatırımın henüz ciroya dönmediğini gösterebilir."),
        ("alacak_hiz", "Alacak devir hızı", "Hasılat / Ortalama ticari alacaklar", lambda c: f"{g('hasilat',c)}/{ort('tic_alacak',c)}", KAT,
         "Yılda kaç kez tahsilat döngüsü tamamlanıyor."),
        ("alacak_gun", "Alacak tahsil süresi (gün)", "365 / Alacak devir hızı", lambda c: f"365*{ort('tic_alacak',c)}/{g('hasilat',c)}", GUN,
         "Uzaması tahsilat sorunu veya agresif satış işareti olabilir."),
        ("stok_hiz", "Stok devir hızı", "Satışların maliyeti / Ortalama stoklar", lambda c: f"{g('smm',c)}/{ort('stok',c)}", KAT,
         "Yılda kaç kez stok yenileniyor."),
        ("stok_gun", "Stokta kalma süresi (gün)", "365 / Stok devir hızı", lambda c: f"365*{ort('stok',c)}/{g('smm',c)}", GUN,
         "Uzaması talep zayıflığı veya eskiyen stok riski olabilir."),
        ("borc_gun", "Ticari borç ödeme süresi (gün)", "365 x Ortalama ticari borçlar / Satışların maliyeti", lambda c: f"365*{ort('tic_borc',c)}/{g('smm',c)}", GUN,
         "Uzun olması tedarikçiye karşı pazarlık gücüne işaret eder."),
        ("ndd", "Nakit dönüşüm süresi (gün)", "Alacak süresi + Stok süresi - Borç süresi", None, GUN,
         "Kısa veya negatifse şirket büyürken nakit tüketmez. Uzaması işletme sermayesi ihtiyacını artırır."),
        ("KARLILIK ORANLARI",),
        ("brut_marj", "Brüt kar marjı", "(Hasılat - SMM) / Hasılat", lambda c: f"({g('hasilat',c)}-{g('smm',c)})/{g('hasilat',c)}", PCT,
         "Fiyatlama gücünün ilk göstergesi. Yıllar içinde istikrarı seviyesinden önemli."),
        ("favok_marj", "FAVÖK marjı", "(Brüt kar - GYG - Pazarlama - ArGe + Amortisman) / Hasılat", lambda c: f"{favok(c)}/{g('hasilat',c)}", PCT,
         "Diğer faaliyet gelir ve giderleri hariç tutulmuştur, böylece kur farkı gibi tek seferlik kalemler marjı bozmaz."),
        ("faal_marj", "Esas faaliyet kar marjı", "Esas faaliyet karı / Hasılat", lambda c: f"{g('faal_kar',c)}/{g('hasilat',c)}", PCT,
         "FAVÖK marjından belirgin farklıysa diğer faaliyet gelir ve giderlerini dipnottan inceleyin."),
        ("net_marj", "Net kar marjı", "Ana ortaklık net karı / Hasılat", lambda c: f"{g('ana_net',c)}/{g('hasilat',c)}", PCT,
         "Faaliyet marjından çok yüksekse karın kaynağı faaliyet dışı olabilir."),
        ("roe", "Özkaynak karlılığı (ROE)", "Ana ortaklık net karı / Ort. ana ortaklık özkaynakları", lambda c: f"{g('ana_net',c)}/{ort('ana_oz',c)}", PCT,
         "Enflasyonun ve faiz oranının üzerinde olması beklenir. Yüksek kaldıraçla şişebilir."),
        ("roa", "Aktif karlılığı (ROA)", "Dönem karı / Ortalama toplam varlıklar", lambda c: f"{g('net_kar',c)}/{ort('tv',c)}", PCT,
         "Kaldıraçtan bağımsız verimlilik ölçüsü."),
        ("eff_vergi", "Efektif vergi oranı", "Vergi gideri / Vergi öncesi kar", lambda c: f"{g('vergi',c)}/{g('vok',c)}", PCT,
         "Yasal oranın çok altındaysa teşvik veya ertelenmiş vergi etkisini dipnotta arayın."),
        ("nopat", "NOPAT (mn TL)", "Esas faaliyet karı x (1 - Efektif vergi oranı)", lambda c: f"{g('faal_kar',c)}*(1-{g('vergi',c)}/{g('vok',c)})", NUM,
         "Vergi sonrası faaliyet karı, ROIC hesabının payı."),
        ("yat_serm", "Yatırılan sermaye (mn TL)", "Toplam özkaynaklar + Net borç", lambda c: f"{g('top_oz',c)}+{netborc(c)}", NUM,
         "Şirketin faaliyetine bağlanmış toplam sermaye."),
        ("roic", "Yatırılan sermaye getirisi (ROIC)", "NOPAT / Ortalama yatırılan sermaye", None, PCT,
         "Sermaye maliyetinin üzerinde kalıcı ROIC değer yaratan şirketi gösterir."),
        ("NAKİT AKIŞI ORANLARI",),
        ("sna", "Serbest nakit akışı (mn TL)", "İşletme nakit akışı - CAPEX", lambda c: f"{g('ocf',c)}-{g('capex',c)}", NUM,
         "Kalıcı negatif SNA borç veya sermaye artırımıyla finanse edilen büyüme demektir."),
        ("sna_marj", "SNA marjı", "SNA / Hasılat", lambda c: f"({g('ocf',c)}-{g('capex',c)})/{g('hasilat',c)}", PCT,
         "Satışların ne kadarının gerçek nakde döndüğü."),
        ("kar_kal", "Kar kalitesi", "İşletme nakit akışı / Dönem karı", lambda c: f"{g('ocf',c)}/{g('net_kar',c)}", KAT,
         "1 civarı veya üzeri sağlıklı. Yıllarca 1 altında kalması tahakkuk ağırlıklı kar sinyali."),
        ("capex_amort", "CAPEX / Amortisman", "CAPEX / Amortisman", lambda c: f"{g('capex',c)}/{g('amort',c)}", KAT,
         "1 üzeri büyüme yatırımı, 1 altı varlık tabanının eridiğini gösterebilir."),
        ("capex_ciro", "CAPEX / Hasılat", "CAPEX / Hasılat", lambda c: f"{g('capex',c)}/{g('hasilat',c)}", PCT,
         "Sektörün yatırım yoğunluğunu gösterir."),
        ("tem_sna", "Temettü / SNA", "Ödenen temettü / SNA", lambda c: f"{g('temettu',c)}/({g('ocf',c)}-{g('capex',c)})", PCT,
         "%100 üzeri temettünün borçla veya nakit stokuyla ödendiğini gösterir."),
    ]
    r = 6
    for item in RAS:
        if len(item) == 1:
            bolum(wr, r, item[0], 6)
            r += 1
            continue
        key, ad, form, fn, fmt, yorum = item
        rsatir[key] = r
        wr.cell(row=r, column=1, value=ad).font = f_norm
        fc = wr.cell(row=r, column=2, value=form)
        fc.font, fc.alignment = Font(name=FONT, size=9, color="555555"), sar
        yc = wr.cell(row=r, column=6, value=yorum)
        yc.font, yc.alignment = Font(name=FONT, size=9), sar
        wr.row_dimensions[r].height = 36
        for col in PER_COLS:
            if key == "ndd":
                expr = f"{col}{rsatir['alacak_gun']}+{col}{rsatir['stok_gun']}-{col}{rsatir['borc_gun']}"
            elif key == "roic":
                expr = f"{col}{rsatir['nopat']}/{ortr('yat_serm', col)}"
            else:
                expr = fn(col)
            c = wr[f"{col}{r}"]
            c.value = f'=IFERROR({expr},"")'
            c.font, c.number_format, c.border = f_norm, fmt, kenar
            c.alignment = Alignment(horizontal="right", vertical="top")
        for col in "ABF":
            wr[f"{col}{r}"].border = kenar
        r += 1
    wr.freeze_panes = "C6"
    wr.cell(row=r + 1, column=1, value="Ortalama bakiye kullanan oranlarda Dönem 1 için dönem sonu bakiyesi kullanılır, bu yüzden ilk sütun bir sonrakilerle tam kıyaslanabilir değildir.").font = f_alt

    # ---------------- Kontrol ----------------
    wk = wb.create_sheet("Kontrol")
    wk.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", [58, 4, 16, 16, 16]):
        wk.column_dimensions[col].width = w
    wk["A1"] = "Girdi Kontrolleri"
    wk["A1"].font = f_baslik
    kafa(wk, 3, ["Kontrol", "", "Dönem 1", "Dönem 2", "Dönem 3"])
    kontroller = [
        ("Toplam varlıklar = KV + UV yükümlülükler + Özkaynaklar",
         lambda c: f'IF({g("tv",c)}="","",IF(ABS({g("tv",c)}-{g("kvyk",c)}-{g("uvyk",c)}-{g("top_oz",c)})<=0.01*{g("tv",c)},"TAMAM","UYARI"))'),
        ("Dönen varlıklar toplam varlıklardan küçük",
         lambda c: f'IF({g("tv",c)}="","",IF({g("donen",c)}<{g("tv",c)},"TAMAM","UYARI"))'),
        ("Ana ortaklık özkaynağı toplam özkaynaktan büyük değil",
         lambda c: f'IF({g("top_oz",c)}="","",IF({g("ana_oz",c)}<={g("top_oz",c)}+0.5,"TAMAM","UYARI"))'),
        ("Giderler pozitif girilmiş (SMM, GYG, pazarlama)",
         lambda c: f'IF({g("hasilat",c)}="","",IF(AND({g("smm",c)}>=0,{g("gyg",c)}>=0,{g("pazarlama",c)}>=0),"TAMAM","UYARI"))'),
        ("CAPEX ve temettü pozitif girilmiş",
         lambda c: f'IF({g("ocf",c)}="","",IF(AND({g("capex",c)}>=0,{g("temettu",c)}>=0),"TAMAM","UYARI"))'),
    ]
    for i, (ad, fn) in enumerate(kontroller):
        rr = 4 + i
        wk.cell(row=rr, column=1, value=ad).font = f_norm
        for col in PER_COLS:
            c = wk[f"{col}{rr}"]
            c.value = "=" + fn(col)
            c.font, c.alignment, c.border = f_kalin, orta, kenar
    wk.cell(row=10, column=1, value="Bilanço denkliğinde yüzde 1 sapma toleransı kullanılır, yuvarlamadan kaynaklı farklar UYARI vermez.").font = f_alt
    wb.save(path)
    return satir
