from openpyxl import Workbook
from ortak import *

GIRDILER = [
    ("BİLANÇO", None, None, None),
    ("donen", "Toplam dönen varlıklar", "Bilanço, Toplam Dönen Varlıklar", "Altman X1, Beneish AQI"),
    ("tv", "Toplam varlıklar", "Bilanço, Toplam Varlıklar", "Altman tümü, Beneish AQI TATA LVGI"),
    ("kvyk", "Toplam kısa vadeli yükümlülükler", "Bilanço, Toplam Kısa Vadeli Yükümlülükler", "Altman X1 X4, Beneish LVGI"),
    ("uvyk", "Toplam uzun vadeli yükümlülükler", "Bilanço, Toplam Uzun Vadeli Yükümlülükler", "Altman X4"),
    ("uv_fin", "Uzun vadeli finansal borçlar", "Bilanço, UV Yükümlülükler, Finansal Borçlar", "Beneish LVGI"),
    ("top_oz", "Toplam özkaynaklar (defter değeri)", "Bilanço, Toplam Özkaynaklar", "Altman Z'' X4"),
    ("gecmis", "Geçmiş yıllar karları (zararları)", "Bilanço, Özkaynaklar, Geçmiş Yıllar Kar/Zararları", "Altman X2"),
    ("donem_bil", "Dönem net karı (bilançodaki)", "Bilanço, Özkaynaklar, Dönem Net Kar/Zararı", "Altman X2"),
    ("tic_alacak", "Ticari alacaklar", "Bilanço, Dönen Varlıklar, Ticari Alacaklar", "Beneish DSRI"),
    ("mdv", "Maddi duran varlıklar (net)", "Bilanço, Duran Varlıklar, Maddi Duran Varlıklar", "Beneish AQI DEPI"),
    ("fin_yat", "Finansal yatırımlar (KV + UV)", "Bilanço, Finansal Yatırımlar (dönen ve duran toplamı)", "Beneish AQI"),
    ("GELİR TABLOSU (giderleri pozitif girin)", None, None, None),
    ("hasilat", "Hasılat", "Gelir Tablosu, Hasılat", "Altman X5, Beneish DSRI GMI SGI SGAI"),
    ("smm", "Satışların maliyeti", "Gelir Tablosu, Satışların Maliyeti (-)", "Beneish GMI"),
    ("gyg", "Genel yönetim giderleri", "Gelir Tablosu, Genel Yönetim Giderleri (-)", "Beneish SGAI"),
    ("pazarlama", "Pazarlama, satış ve dağıtım giderleri", "Gelir Tablosu, Pazarlama Giderleri (-)", "Beneish SGAI"),
    ("fvok", "Faiz ve vergi öncesi kar (FVÖK, EBIT)", "Gelir Tablosu, Finansman Geliri (Gideri) Öncesi Faaliyet Karı", "Altman X3"),
    ("surd_kar", "Sürdürülen faaliyetler dönem karı", "Gelir Tablosu, Sürdürülen Faaliyetler Dönem Karı", "Beneish TATA"),
    ("NAKİT AKIŞ TABLOSU", None, None, None),
    ("amort", "Amortisman ve itfa payları", "Nakit Akış Tablosu, Amortisman ve İtfa Gideri ile İlgili Düzeltmeler", "Beneish DEPI"),
    ("ocf", "İşletme faaliyetlerinden nakit akışı", "Nakit Akış Tablosu, A. İşletme Faaliyetlerinden Nakit Akışları", "Beneish TATA"),
    ("PİYASA VERİSİ", None, None, None),
    ("pd_nom", "Yıl sonu piyasa değeri (nominal)", "Bilanço tarihindeki kapanış fiyatı x Pay adedi", "Altman X4"),
    ("pd_kat", "Piyasa değeri endeksleme katsayısı", "TMS 29 uygulanmıyorsa 1. Uygulanıyorsa tabloların satın alma gücü tarihine taşıyan TÜFE katsayısı", "Altman X4"),
]
P = ["C", "D", "E"]


def build(path, donemler=None, veri=None, notlar=None, baslik_ek=""):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Nasıl Kullanılır"
    ws0.sheet_view.showGridLines = False
    ws0.column_dimensions["A"].width = 4
    ws0.column_dimensions["B"].width = 100
    ws0.column_dimensions["C"].width = 40
    ws0["B2"] = "Altman Z ve Beneish M Hesap Şablonu" + baslik_ek
    ws0["B2"].font = f_baslik
    ws0["B3"] = "Bilanço Okuma Eğitimi, ek materyal 5"
    ws0["B3"].font = f_alt
    metin = [
        "1. Girdi sayfasında her satırın yanında rakamı tablonun hangi kaleminden alacağınız ve hangi bileşende kullanıldığı yazıyor.",
        "2. En az iki ardışık yıl girin. Beneish M her yılı bir önceki yılla kıyasladığı için ilk dönem sütununda hesaplanmaz.",
        "3. Tutarları aynı birimde girin (şablon milyon TL varsayıyor). Giderleri pozitif girin.",
        "4. Altman Z orijinal modeli halka açık sanayi şirketleri için geliştirildi. Hizmet, ticaret ve gelişmekte olan piyasa şirketlerinde Z'' modeline de bakın.",
        "5. Beneish M 8 değişkenli modeldir. -1,78 eşiğinin üzeri kar manipülasyonu olasılığının yüksek olduğu bölgedir. Bazı kaynaklar daha temkinli -2,22 eşiğini kullanır.",
        "6. Bu iki skor tarama aracıdır, hüküm değildir. Eşiği aşan şirkette hangi endeksin skoru yükselttiğine bakın ve ilgili dipnotu okuyun.",
        "7. TMS 29 uygulayan şirketlerde bilanço rakamları raporlama tarihinin satın alma gücüyle ifade edilir. Piyasa değerini de aynı alım gücüne taşımak için endeksleme katsayısını kullanın, aksi halde X4 bozulur.",
        "8. Bankalar, sigorta şirketleri ve holdingler için bu modeller uygun değildir.",
    ]
    for i, a in enumerate(metin):
        c = ws0.cell(row=5 + i, column=2, value=a)
        c.font, c.alignment = f_norm, sar
        ws0.row_dimensions[5 + i].height = 30
    lejant(ws0, 15, col=2)

    # Girdi
    ws = wb.create_sheet("Girdi")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [40, 56, 14, 14, 14, 3, 30]):
        ws.column_dimensions[col].width = w
    ws["A1"] = "Girdi Tablosu (milyon TL)" + baslik_ek
    ws["A1"].font = f_baslik
    ws["A2"] = "Sarı hücreleri doldurun. En eski dönem solda."
    ws["A2"].font = f_alt
    kafa(ws, 4, ["Girdi", "Finansal tablo kalemi", "Dönem 1", "Dönem 2", "Dönem 3"])
    kafa(ws, 4, ["Kullanıldığı yer"], col=7)
    ws["A5"] = "Dönem adı"
    ws["A5"].font = f_kalin
    for i, col in enumerate(P):
        c = ws[f"{col}5"]
        c.value = donemler[i] if donemler else None
        c.font, c.fill, c.border, c.alignment = f_girdi, fill_girdi, kenar, orta
    S = {}
    r = 6
    for key, ad, kalem, yer in GIRDILER:
        if ad is None:
            bolum(ws, r, key, 5)
            r += 1
            continue
        S[key] = r
        ws.cell(row=r, column=1, value=ad).font = f_norm
        k = ws.cell(row=r, column=2, value=kalem)
        k.font, k.alignment = Font(name=FONT, size=9, color="555555"), sar
        ws.cell(row=r, column=7, value=yer).font = Font(name=FONT, size=9, color="555555")
        for i, col in enumerate(P):
            c = ws[f"{col}{r}"]
            if veri:
                c.value = veri[key][i]
            elif key == "pd_kat":
                c.value = 1
            c.font, c.fill, c.border = f_girdi, fill_girdi, kenar
            c.number_format = "0.0000" if key == "pd_kat" else NUM
            if veri and notlar and key in notlar:
                c.comment = Comment(notlar[key], "Eğitim")
        ws.cell(row=r, column=1).border = kenar
        ws.cell(row=r, column=2).border = kenar
        r += 1
    # türetilen
    bolum(ws, r, "TÜRETİLEN KALEMLER (formül)", 5)
    r += 1
    turet = [
        ("pd", "Piyasa değeri (tablo alım gücüyle)", lambda c: f"{c}{S['pd_nom']}*{c}{S['pd_kat']}"),
        ("top_yuk", "Toplam yükümlülükler", lambda c: f"{c}{S['kvyk']}+{c}{S['uvyk']}"),
        ("dagitilmamis", "Dağıtılmamış karlar", lambda c: f"{c}{S['gecmis']}+{c}{S['donem_bil']}"),
        ("nis", "Net çalışma sermayesi", lambda c: f"{c}{S['donen']}-{c}{S['kvyk']}"),
        ("sga", "Satış, genel ve yönetim giderleri", lambda c: f"{c}{S['gyg']}+{c}{S['pazarlama']}"),
    ]
    for key, ad, fn in turet:
        S[key] = r
        ws.cell(row=r, column=1, value=ad).font = f_norm
        for col in P:
            c = ws[f"{col}{r}"]
            c.value = f'=IF({col}{S["tv"]}="","",{fn(col)})'
            c.font, c.number_format, c.border = f_norm, NUM, kenar
        r += 1
    ws.freeze_panes = "C6"
    if veri and "_kaynak" in veri:
        ws.cell(row=r + 1, column=1, value=veri["_kaynak"]).font = f_alt

    def g(k, c):
        return f"Girdi!{c}{S[k]}"

    # Altman
    wa = wb.create_sheet("Altman Z")
    wa.sheet_view.showGridLines = False
    for col, w in zip("ABCDEF", [36, 46, 13, 13, 13, 10]):
        wa.column_dimensions[col].width = w
    wa["A1"] = "Altman Z Skoru" + baslik_ek
    wa["A1"].font = f_baslik
    wa["A2"] = "Tüm hücreler formüldür."
    wa["A2"].font = f_alt
    kafa(wa, 4, ["Bileşen", "Tanım", "Dönem 1", "Dönem 2", "Dönem 3", "Katsayı"])
    for col in P:
        wa[f"{col}5"] = f'=IF(Girdi!{col}5="","",Girdi!{col}5)'
        wa[f"{col}5"].font, wa[f"{col}5"].alignment = Font(name=FONT, size=10, bold=True, color="008000"), orta
    A = {}
    bolum(wa, 6, "ORİJİNAL MODEL (1968, halka açık sanayi şirketi)", 6)
    xs = [
        ("x1", "X1", "Net çalışma sermayesi / Toplam varlıklar", lambda c: f"{g('nis',c)}/{g('tv',c)}", 1.2),
        ("x2", "X2", "Dağıtılmamış karlar / Toplam varlıklar", lambda c: f"{g('dagitilmamis',c)}/{g('tv',c)}", 1.4),
        ("x3", "X3", "FVÖK / Toplam varlıklar", lambda c: f"{g('fvok',c)}/{g('tv',c)}", 3.3),
        ("x4", "X4", "Piyasa değeri / Toplam yükümlülükler", lambda c: f"{g('pd',c)}/{g('top_yuk',c)}", 0.6),
        ("x5", "X5", "Hasılat / Toplam varlıklar", lambda c: f"{g('hasilat',c)}/{g('tv',c)}", 1.0),
    ]
    r = 7
    for key, ad, tanim, fn, kat in xs:
        A[key] = r
        wa.cell(row=r, column=1, value=ad).font = f_norm
        wa.cell(row=r, column=2, value=tanim).font = Font(name=FONT, size=9, color="555555")
        for col in P:
            c = wa[f"{col}{r}"]
            c.value = f'=IFERROR({fn(col)},"")'
            c.number_format, c.font, c.border = DEC, f_norm, kenar
        k = wa.cell(row=r, column=6, value=kat)
        k.font, k.number_format = f_girdi, "0.00"
        r += 1
    A["z"] = r
    wa.cell(row=r, column=1, value="Z skoru").font = f_kalin
    wa.cell(row=r, column=2, value="1,2 X1 + 1,4 X2 + 3,3 X3 + 0,6 X4 + 1,0 X5").font = Font(name=FONT, size=9, color="555555")
    for col in P:
        c = wa[f"{col}{r}"]
        c.value = f'=IFERROR(SUMPRODUCT({col}{A["x1"]}:{col}{A["x5"]},$F${A["x1"]}:$F${A["x5"]}),"")'
        c.value = f'=IF({col}{A["x1"]}="","",SUMPRODUCT({col}{A["x1"]}:{col}{A["x5"]},$F${A["x1"]}:$F${A["x5"]}))'
        c.font, c.fill, c.number_format, c.border = f_kalin, fill_sonuc, "0.00", kenar
    r += 1
    A["bolge"] = r
    wa.cell(row=r, column=1, value="Bölge").font = f_kalin
    wa.cell(row=r, column=2, value="2,99 üzeri güvenli, 1,81 ile 2,99 arası gri, 1,81 altı tehlike").font = Font(name=FONT, size=9, color="555555")
    for col in P:
        z = f"{col}{A['z']}"
        c = wa[f"{col}{r}"]
        c.value = f'=IF({z}="","",IF({z}>2.99,"Güvenli",IF({z}>=1.81,"Gri","Tehlike")))'
        c.font, c.alignment, c.border = f_kalin, orta, kenar
    r += 2
    bolum(wa, r, "Z'' MODELİ (1995, üretim dışı ve gelişmekte olan piyasa şirketleri)", 6)
    r += 1
    xs2 = [
        ("y1", "X1", "Net çalışma sermayesi / Toplam varlıklar", f"x1", 6.56),
        ("y2", "X2", "Dağıtılmamış karlar / Toplam varlıklar", f"x2", 3.26),
        ("y3", "X3", "FVÖK / Toplam varlıklar", f"x3", 6.72),
        ("y4", "X4", "Özkaynak defter değeri / Toplam yükümlülükler", None, 1.05),
    ]
    for key, ad, tanim, ref, kat in xs2:
        A[key] = r
        wa.cell(row=r, column=1, value=ad).font = f_norm
        wa.cell(row=r, column=2, value=tanim).font = Font(name=FONT, size=9, color="555555")
        for col in P:
            c = wa[f"{col}{r}"]
            c.value = f"={col}{A[ref]}" if ref else f'=IFERROR({g("top_oz",col)}/{g("top_yuk",col)},"")'
            c.number_format, c.font, c.border = DEC, f_norm, kenar
        k = wa.cell(row=r, column=6, value=kat)
        k.font, k.number_format = f_girdi, "0.00"
        r += 1
    A["z2"] = r
    wa.cell(row=r, column=1, value="Z'' skoru").font = f_kalin
    wa.cell(row=r, column=2, value="6,56 X1 + 3,26 X2 + 6,72 X3 + 1,05 X4").font = Font(name=FONT, size=9, color="555555")
    for col in P:
        c = wa[f"{col}{r}"]
        c.value = f'=IF({col}{A["y1"]}="","",SUMPRODUCT({col}{A["y1"]}:{col}{A["y4"]},$F${A["y1"]}:$F${A["y4"]}))'
        c.font, c.fill, c.number_format, c.border = f_kalin, fill_sonuc, "0.00", kenar
    r += 1
    wa.cell(row=r, column=1, value="Bölge").font = f_kalin
    wa.cell(row=r, column=2, value="2,60 üzeri güvenli, 1,10 ile 2,60 arası gri, 1,10 altı tehlike").font = Font(name=FONT, size=9, color="555555")
    for col in P:
        z = f"{col}{A['z2']}"
        c = wa[f"{col}{r}"]
        c.value = f'=IF({z}="","",IF({z}>2.6,"Güvenli",IF({z}>=1.1,"Gri","Tehlike")))'
        c.font, c.alignment, c.border = f_kalin, orta, kenar
    r += 2
    for t in ["Katsayılar Altman (1968) ve Altman (1995) çalışmalarından alınmıştır.",
              "Dağıtılmamış karlar burada Geçmiş yıllar karları ile Dönem net karının toplamı olarak alınır. Kısıtlanmış yedekleri eklemek isteyen Girdi sayfasında Geçmiş yıllar satırına ekleyebilir.",
              "TMS 29 altında özkaynak kalemleri yeniden ifade edildiği için X2 enflasyon öncesi yıllara göre farklı davranabilir. Trendi aynı muhasebe rejimi içinde kıyaslayın."]:
        wa.cell(row=r, column=1, value=t).font = f_alt
        r += 1

    # Beneish
    wbn = wb.create_sheet("Beneish M")
    wbn.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [12, 54, 13, 13, 13, 10, 44]):
        wbn.column_dimensions[col].width = w
    wbn["A1"] = "Beneish M Skoru (8 değişkenli)" + baslik_ek
    wbn["A1"].font = f_baslik
    wbn["A2"] = "Her sütun bir önceki dönemle kıyaslanır, Dönem 1 sütunu bu yüzden boş kalır."
    wbn["A2"].font = f_alt
    kafa(wbn, 4, ["Endeks", "Formül", "Dönem 1", "Dönem 2", "Dönem 3", "Katsayı", "1'in üzerindeyse ne düşünmeli"])
    for col in P:
        wbn[f"{col}5"] = f'=IF(Girdi!{col}5="","",Girdi!{col}5)'
        wbn[f"{col}5"].font, wbn[f"{col}5"].alignment = Font(name=FONT, size=10, bold=True, color="008000"), orta

    def bn(fn):
        def f(c, p):
            return fn(c, p)
        return f
    idx = [
        ("DSRI", "(Alacak t / Satış t) / (Alacak t-1 / Satış t-1)",
         lambda c, p: f"({g('tic_alacak',c)}/{g('hasilat',c)})/({g('tic_alacak',p)}/{g('hasilat',p)})", 0.920,
         "Alacaklar satışlardan hızlı büyüyor, gelir öne çekilmiş olabilir."),
        ("GMI", "Brüt marj t-1 / Brüt marj t",
         lambda c, p: f"(({g('hasilat',p)}-{g('smm',p)})/{g('hasilat',p)})/(({g('hasilat',c)}-{g('smm',c)})/{g('hasilat',c)})", 0.528,
         "Brüt marj daralıyor, yönetim üzerinde baskı artar."),
        ("AQI", "[1 - (Dönen + MDV + Fin. yat.) / TV] t / aynı oran t-1",
         lambda c, p: f"(1-({g('donen',c)}+{g('mdv',c)}+{g('fin_yat',c)})/{g('tv',c)})/(1-({g('donen',p)}+{g('mdv',p)}+{g('fin_yat',p)})/{g('tv',p)})", 0.404,
         "Somut olmayan, aktifleştirilmiş varlıkların payı artıyor."),
        ("SGI", "Satış t / Satış t-1",
         lambda c, p: f"{g('hasilat',c)}/{g('hasilat',p)}", 0.892,
         "Hızlı büyüme manipülasyon baskısı yaratabilir. Enflasyonda reel tablolara bakın."),
        ("DEPI", "[Amort / (Amort + MDV)] t-1 / aynı oran t",
         lambda c, p: f"({g('amort',p)}/({g('amort',p)}+{g('mdv',p)}))/({g('amort',c)}/({g('amort',c)}+{g('mdv',c)}))", 0.115,
         "Amortisman oranı düşmüş, faydalı ömürler uzatılmış olabilir."),
        ("SGAI", "(SGA t / Satış t) / (SGA t-1 / Satış t-1)",
         lambda c, p: f"({g('sga',c)}/{g('hasilat',c)})/({g('sga',p)}/{g('hasilat',p)})", -0.172,
         "Faaliyet giderleri satışa göre artıyor, verimlilik düşüyor."),
        ("TATA", "(Sürdürülen faaliyet karı t - İşletme nakit akışı t) / TV t",
         lambda c, p: f"({g('surd_kar',c)}-{g('ocf',c)})/{g('tv',c)}", 4.679,
         "Pozitif ve yüksekse kar nakitle desteklenmiyor, tahakkuk ağırlıklı."),
        ("LVGI", "[(KVYK + UV fin. borç) / TV] t / aynı oran t-1",
         lambda c, p: f"(({g('kvyk',c)}+{g('uv_fin',c)})/{g('tv',c)})/(({g('kvyk',p)}+{g('uv_fin',p)})/{g('tv',p)})", -0.327,
         "Kaldıraç artıyor, borç sözleşmesi baskısı oluşabilir."),
    ]
    r = 6
    B = {}
    for ad, form, fn, kat, yorum in idx:
        B[ad] = r
        wbn.cell(row=r, column=1, value=ad).font = f_kalin
        wbn.cell(row=r, column=2, value=form).font = Font(name=FONT, size=9, color="555555")
        for i, col in enumerate(P):
            c = wbn[f"{col}{r}"]
            if i == 0:
                c.value = "-"
                c.alignment = orta
            else:
                p = P[i - 1]
                c.value = f'=IF(OR(Girdi!{p}{S["tv"]}="",Girdi!{col}{S["tv"]}=""),"",IFERROR({fn(col, p)},""))'
                c.number_format = DEC
            c.font, c.border = f_norm, kenar
        k = wbn.cell(row=r, column=6, value=kat)
        k.font, k.number_format = f_girdi, "0.000"
        y = wbn.cell(row=r, column=7, value=yorum)
        y.font, y.alignment = Font(name=FONT, size=9), sar
        wbn.row_dimensions[r].height = 28
        r += 1
    wbn.cell(row=r, column=1, value="Sabit").font = f_norm
    sabit = wbn.cell(row=r, column=6, value=-4.84)
    sabit.font, sabit.number_format = f_girdi, "0.000"
    B["sabit"] = r
    r += 1
    B["m"] = r
    wbn.cell(row=r, column=1, value="M skoru").font = f_kalin
    wbn.cell(row=r, column=2, value="-4,84 + 0,920 DSRI + 0,528 GMI + 0,404 AQI + 0,892 SGI + 0,115 DEPI - 0,172 SGAI + 4,679 TATA - 0,327 LVGI").font = Font(name=FONT, size=9, color="555555")
    for i, col in enumerate(P):
        c = wbn[f"{col}{r}"]
        if i == 0:
            c.value = "-"
            c.alignment = orta
        else:
            c.value = (f'=IF(COUNT({col}{B["DSRI"]}:{col}{B["LVGI"]})<8,"",'
                       f'$F${B["sabit"]}+SUMPRODUCT({col}{B["DSRI"]}:{col}{B["LVGI"]},$F${B["DSRI"]}:$F${B["LVGI"]}))')
            c.number_format = "0.00"
        c.font, c.fill, c.border = f_kalin, fill_sonuc, kenar
    r += 1
    wbn.cell(row=r, column=1, value="Eşik").font = f_norm
    e = wbn.cell(row=r, column=6, value=-1.78)
    e.font, e.number_format = f_girdi, "0.00"
    B["esik"] = r
    wbn.cell(row=r, column=2, value="Beneish (1999) 8 değişkenli model eşiği. Temkinli kullanım için -2,22 yazabilirsiniz.").font = Font(name=FONT, size=9, color="555555")
    r += 1
    wbn.cell(row=r, column=1, value="Sonuç").font = f_kalin
    for i, col in enumerate(P):
        c = wbn[f"{col}{r}"]
        if i == 0:
            c.value = "-"
        else:
            m = f"{col}{B['m']}"
            c.value = f'=IF({m}="","",IF({m}>$F${B["esik"]},"Riskli bölge","Eşiğin altında"))'
        c.font, c.alignment, c.border = f_kalin, orta, kenar
    r += 2
    for t in ["Katsayılar Beneish (1999), The Detection of Earnings Manipulation, Financial Analysts Journal çalışmasından alınmıştır.",
              "Model ABD şirketleri verisiyle kalibre edilmiştir. Yüksek enflasyon ve kur oynaklığı SGI, GMI ve DSRI endekslerini manipülasyon olmadan da şişirebilir.",
              "Sonucu tek başına kullanmayın. Eşiği aşan durumda en yüksek katkıyı yapan endeksi bulun ve Kırmızı Bayrak Kontrol Listesi ile birlikte değerlendirin."]:
        wbn.cell(row=r, column=1, value=t).font = f_alt
        r += 1
    wb.save(path)
